import shutil
import subprocess
from pathlib import Path

import docx
import pytest

from cvflow.onboarding import parse
from cvflow.onboarding.parse import ResumeParseError, parse_resume, parse_resumes

MASTER = Path(__file__).parent / "fixtures" / "resume" / "master.tex"
needs_tectonic = pytest.mark.skipif(shutil.which("tectonic") is None, reason="no tectonic")


def _compile(tmp_path: Path, tex: str) -> Path:
    src = tmp_path / "doc.tex"
    src.write_text(tex)
    subprocess.run(["tectonic", str(src)], check=True, capture_output=True)  # noqa: S603,S607
    return tmp_path / "doc.pdf"


def test_docx_paragraphs_and_table(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_heading("Jordan Lee", 0)
    d.add_paragraph("Built reliable systems at scale.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Orchestration"
    t.cell(0, 1).text = "Kubernetes"
    path = tmp_path / "r.docx"
    d.save(str(path))
    res = parse_resume(path)
    assert res.kind == "docx"
    assert "Jordan Lee" in res.text
    assert "Built reliable systems" in res.text
    assert "Orchestration | Kubernetes" in res.text


@needs_tectonic
def test_pdf_text(tmp_path: Path) -> None:
    pdf = _compile(
        tmp_path,
        "\\documentclass{article}\\begin{document}"
        "Jordan Lee is a platform engineer with Kubernetes and Terraform experience."
        "\\end{document}",
    )
    res = parse_resume(pdf)
    assert res.kind == "pdf"
    assert "Kubernetes" in res.text


@needs_tectonic
def test_blank_pdf_raises(tmp_path: Path) -> None:
    pdf = _compile(
        tmp_path,
        "\\documentclass{article}\\pagestyle{empty}\\begin{document}\\mbox{}\\end{document}",
    )
    with pytest.raises(ResumeParseError, match="scanned"):
        parse_resume(pdf)


def test_txt_passthrough_and_multi(tmp_path: Path) -> None:
    a = tmp_path / "a.txt"
    a.write_text("alpha")
    b = tmp_path / "b.md"
    b.write_text("beta")
    assert parse_resume(a).text == "alpha"
    assert parse_resume(a).kind == "text"
    both = parse_resumes([a, b])
    assert both.kind == "mixed"
    assert "\n\n=== FILE: b.md ===\n\nbeta" in both.text


def test_unknown_extension(tmp_path: Path) -> None:
    f = tmp_path / "r.rtf"
    f.write_text("x")
    with pytest.raises(ResumeParseError):
        parse_resume(f)


def test_oversize(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    f = tmp_path / "big.txt"
    f.write_text("x")
    monkeypatch.setattr(parse, "MAX_BYTES", 0)
    with pytest.raises(ResumeParseError, match="limit"):
        parse_resume(f)
